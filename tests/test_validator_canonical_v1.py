from __future__ import annotations

import json
import subprocess
import sys
from pathlib import Path

from openpyxl import Workbook


ROOT = Path(__file__).resolve().parents[1]
VALIDATOR_ENTRY = ROOT / "validate_recipe_excel.py"
TEMPLATE = ROOT / "templates" / "recipe_excel_canonical_template_v1.xlsx"
INVALID_DIR = ROOT / "templates" / "invalid_samples"


def run_validator(path: Path) -> tuple[int, dict]:
    proc = subprocess.run(
        [sys.executable, str(VALIDATOR_ENTRY), str(path), "--json"],
        capture_output=True,
        text=True,
        check=False,
        cwd=str(ROOT),
    )
    payload = json.loads(proc.stdout)
    return proc.returncode, payload


def test_canonical_template_passes_all_phases() -> None:
    rc, payload = run_validator(TEMPLATE)
    assert rc == 0
    assert payload["status"] == "passed"
    assert payload["summary"]["error_count"] == 0
    assert payload["phase_execution"]["executed_phases"] == [
        "PHASE_1_STRUCTURE",
        "PHASE_2_VALUE",
        "PHASE_3_UNIQUENESS",
        "PHASE_4_REFERENCE",
    ]
    assert payload["phase_execution"]["skipped_phases"] == []
    assert payload["first_failure"] is None


def test_invalid_sheet_name_is_structure_error_and_stops_after_phase1() -> None:
    rc, payload = run_validator(INVALID_DIR / "invalid_sheet_name.xlsx")
    assert rc == 1
    assert payload["first_failure"]["error_class"] == "STRUCTURE_ERROR"
    assert payload["phase_execution"]["executed_phases"] == ["PHASE_1_STRUCTURE"]
    assert payload["phase_execution"]["skipped_phases"] == [
        "PHASE_2_VALUE",
        "PHASE_3_UNIQUENESS",
        "PHASE_4_REFERENCE",
    ]
    assert all(e["error_class"] == "STRUCTURE_ERROR" for e in payload["all_errors"])


def test_invalid_missing_recipe_column_is_structure_error() -> None:
    rc, payload = run_validator(INVALID_DIR / "invalid_missing_recipe_column.xlsx")
    assert rc == 1
    assert payload["first_failure"]["error_class"] == "STRUCTURE_ERROR"
    assert payload["phase_execution"]["executed_phases"] == ["PHASE_1_STRUCTURE"]
    assert all(e["error_class"] == "STRUCTURE_ERROR" for e in payload["all_errors"])


def test_value_error_samples_first_failure_value_error() -> None:
    targets = [
        "invalid_bad_category_code.xlsx",
        "invalid_bad_tags.xlsx",
        "invalid_servings_not_integer.xlsx",
        "invalid_bad_unit.xlsx",
        "invalid_bad_process_code.xlsx",
        "invalid_missing_food_id.xlsx",
    ]
    for name in targets:
        rc, payload = run_validator(INVALID_DIR / name)
        assert rc == 1, name
        assert payload["first_failure"]["error_class"] == "VALUE_ERROR", name
        assert payload["phase_execution"]["executed_phases"] == [
            "PHASE_1_STRUCTURE",
            "PHASE_2_VALUE",
            "PHASE_3_UNIQUENESS",
            "PHASE_4_REFERENCE",
        ], name


def test_uniqueness_error_samples_first_failure_uniqueness_error() -> None:
    targets = [
        "invalid_duplicate_ingredient_no.xlsx",
        "invalid_step_no_duplicate.xlsx",
    ]
    for name in targets:
        rc, payload = run_validator(INVALID_DIR / name)
        assert rc == 1, name
        assert payload["first_failure"]["error_class"] == "UNIQUENESS_ERROR", name


def test_batch_cook_is_invalid_tag() -> None:
    _, payload = run_validator(INVALID_DIR / "invalid_bad_tags.xlsx")
    tag_errors = [e for e in payload["all_errors"] if e["error_code"] == "INVALID_TAG_VALUE"]
    assert tag_errors
    assert any(e["details"].get("tag") == "batch_cook" for e in tag_errors)


def test_square_ml_character_is_invalid_unit() -> None:
    _, payload = run_validator(INVALID_DIR / "invalid_bad_unit.xlsx")
    unit_errors = [e for e in payload["all_errors"] if e["error_code"] == "INVALID_UNIT_CHARACTER"]
    assert unit_errors
    assert unit_errors[0]["details"]["value"] == "㎖"


def test_first_failure_and_all_errors_are_both_reported(tmp_path: Path) -> None:
    out = tmp_path / "multi_error.xlsx"
    wb = Workbook()

    ws_r = wb.active
    ws_r.title = "Recipes"
    ws_i = wb.create_sheet("Ingredients")
    ws_s = wb.create_sheet("Steps")

    ws_r.append(["Recipe_ID", "Recipe_Name", "Category_Code", "Servings", "Tags"])
    ws_i.append(["Recipe_ID", "Ingredient_No", "Food_ID", "Ingredient_Name", "Amount_Value", "Unit", "Process_Code"])
    ws_s.append(["Recipe_ID", "Step_No", "Instruction"])

    ws_r.append(["R001", "A", "rice", 0, "dinner"])  # VALUE_ERROR: Servings invalid
    ws_i.append(["R001", 1, "F001", "water", 100, "ml", "RAW"])
    ws_i.append(["R001", 1, "F002", "salt", 2, "g", "RAW"])  # UNIQUENESS_ERROR
    ws_i.append(["R999", 2, "F003", "sugar", 3, "g", "RAW"])  # REFERENCE_ERROR
    ws_s.append(["R001", 1, "Do it"])

    wb.save(out)

    rc, payload = run_validator(out)
    assert rc == 1
    assert payload["first_failure"]["error_class"] == "VALUE_ERROR"
    classes = {e["error_class"] for e in payload["all_errors"]}
    assert "VALUE_ERROR" in classes
    assert "UNIQUENESS_ERROR" in classes
    assert "REFERENCE_ERROR" in classes
    first_failure_count = sum(1 for e in payload["all_errors"] if e["is_first_failure"])
    assert first_failure_count == 1
